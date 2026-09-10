"""Build a linked review deck from preserved boards; never modify approved images.

Requires python-pptx. Run from any directory. Navigation links are internal slides.
This is a storyboard, not a runnable app or authorization to implement workflows.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'visual_baselines' / 'presentation'
OUT.mkdir(parents=True, exist_ok=True)
P = Presentation()
P.slide_width, P.slide_height = Inches(16), Inches(10)
BG, INK, TEAL, MUTED = 'F4F7F7', '192E33', '287A80', '51656B'
groups = [
 ('Acceso y equipo compartido', 'Elegir modalidad por equipo; Workspace y PIN son alternativas, no pasos obligatorios consecutivos.', [
 ('ACC-01','Usuario','Ingresar con credenciales a Workspace.'),('ACC-02','Vendedor','Alternativa: PIN en equipo habilitado. Sólo capacidades de venta.'),('ACC-03','Usuario multirrol','Acceder a las áreas autorizadas sin cambiar de perfil.')]),
 ('Mostrador y venta consultiva', 'Dos recorridos alternativos sobre órdenes de venta. La aprobación depende de la política de Odoo.', [
 ('VEN-01','Vendedor','Buscar una orden existente o iniciar una nueva.'),('VEN-05','Vendedor','Consultar cliente/producto cuando sea necesario.'),('VEN-03','Vendedor de mostrador','Capturar líneas rápidamente; búsqueda inline. Derivar a Caja según condición.'),('VEN-04','Vendedor consultivo','Preparar cotización, secciones y condiciones. Solicitar aprobación si corresponde.'),('SUP-01','Supervisor','Resolver solicitud; devolver a Ventas. No todas las órdenes necesitan este paso.')]),
 ('Caja: cobrar y consultar', 'El cajero entra con credenciales. Cobrar una venta y cobrar cartera son recorridos diferentes.', [
 ('CAJ-09-v2','Cajero','Reconocer punto, sesión y acciones disponibles.'),('PENDIENTES-RESULTADO-COBRO-v1','Cajero','Abrir el pendiente recibido de Ventas. Lámina previamente aprobada.'),('CAJ-11','Cajero','Revisar pago, saldo y resultado; conservar identidad ante interrupción.'),('CAJ-03','Cajero','Alternativa: aplicar un cobro a facturas de cartera.'),('CAJ-02','Cajero / supervisor','Consultar registros y comprobantes, sin volver a cobrar.')]),
 ('Caja: operaciones auxiliares y turno', 'Operaciones independientes desde Caja: no ejecutar una detrás de otra. Reutilizar procesos reales de Odoo.', [
 ('CAJ-04-v2','Cajero','Retención: consultar clave, revisar documentos, registrar.'),('CAJ-05-v2','Cajero','Anticipo: capturar medios de pago, procesar y consultar disponible.'),('CAJ-06-v2','Cajero autorizado','Depósito: registrar datos y contabilizar según permisos.'),('CAJ-07-v2','Cajero autorizado','Salida: revisar tipo y documentos antes de aceptar.'),('CAJ-08-v2','Cajero autorizado','Cruce: asignar fuentes a destinos; no constituye nuevo cobro en efectivo.'),('CAJ-10','Cajero / supervisor','Turno: apertura, arqueo, revisión de diferencias y cierre. Estados aún por detallar.')]),
 ('Bodega e inventario normal', 'Preparar y entregar no son equivalentes. Los candados y decisiones proceden de Odoo.', [
 ('BOD-01','Bodeguero','Consultar productos normales y prioridades de trabajo.'),('BOD-02','Bodeguero','Abrir recepción, preparación, entrega o transferencia.'),('BOD-03','Bodeguero','Preparar cantidades, registrar faltantes y revisar entrega autorizada.'),('BOD-04','Responsable de inventario','Recorrido alternativo: contar productos y enviar diferencias a revisión.')]),
 ('Envases propios: custodia y movimiento', 'Workspace independiente, sin sesión de Caja. Contenido, recipiente y presentación son conceptos separados.', [
 ('ENV-01','Responsable de envases','Ver ubicación y custodia consolidada por producto/envase.'),('ENV-FACTURA-v1','Responsable de envases','Origen alternativo A: registrar desde factura; propuesta aprobada.'),('ENV-TOMA-FISICA-v1','Responsable de envases','Origen alternativo B: toma física; propuesta aprobada. No sumar conteos repetidos.'),('ENV-03','Responsable de envases','Consultar entregas, devoluciones y tránsitos por atender.'),('ENV-06','Responsable de envases','Registrar recepción/devolución por múltiples productos y cantidades.'),('ENV-02','Responsable de envases','Consultar historial y documentos vinculados.'),('ENV-07','Usuario autorizado','Recorrido comercial separado: comprar o vender envases mediante documentos de Odoo.')]),
 ('Offline, conflictos y continuidad', 'Estado comercial y estado de sincronización son independientes; un reintento no crea otro documento.', [
 ('SYN-01','Usuario','Preparar catálogos y comprobar cobertura offline.'),('SYN-02','Usuario','Consultar operaciones locales, dependencias y reintentos.'),('SYN-03','Responsable autorizado','Si hay conflicto: comparar y resolver sin borrar el trabajo.'),('OPS-01','Usuario / soporte','Si hay incidente: recuperar acceso o espacio conservando pendientes.')]),
 ('Preferencias y seguimiento', 'Acciones transversales que no sustituyen los permisos ni las reglas empresariales.', [
 ('CFG-01','Usuario / administrador','Configurar apariencia y modalidad del equipo según permisos.'),('NOT-01','Usuario','Abrir la actividad o notificación y volver al documento relacionado.')]),
]

def txt(s, value, x,y,w,h,size=18,color=INK,bold=False):
    box=s.shapes.add_textbox(Inches(x),Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.word_wrap=True
    for a in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,a,0)
    p=tf.paragraphs[0]; p.text=value; p.font.size=Pt(size);p.font.bold=bold
    p.font.name='Aptos';p.font.color.rgb=RGBColor.from_string(color)
    return box

def slide(title, subtitle):
    s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    txt(s,'ORBI ERP  /  REVISIÓN DE RECORRIDOS',.5,.2,13,.3,12,TEAL,True)
    txt(s,title,.5,.65,15,.6,30,INK,True)
    txt(s,subtitle,.5,1.4,15,.65,16,MUTED)
    txt(s,'Storyboard de revisión · datos ficticios · no es una app funcional',.5,9.62,13,.23,10,MUTED)
    txt(s,str(len(P.slides)),15,9.62,.5,.23,10,MUTED)
    return s

def button(s,label,x,y,w,target):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(.46))
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(TEAL);sh.line.fill.background()
    sh.text=label
    for p in sh.text_frame.paragraphs:
        p.font.size=Pt(13);p.font.color.rgb=RGBColor(255,255,255);p.font.name='Aptos'
    sh.click_action.target_slide=target

index=slide('Los flujos de Orbi, pantalla por pantalla','Pulsa un recorrido. Los botones Índice / Anterior / Siguiente funcionan en modo presentación.')
sections=[]; detail_count=0
for title, note, steps in groups:
    intro=slide(title,note); sections.append((title,intro))
    cards=[]
    for n,(sid,actor,action) in enumerate(steps):
        approved=sid.endswith('-v1')
        path=ROOT/'visual_baselines'/('approved' if approved else 'approved/round-02')/(sid+'.png')
        if not path.is_file(): raise FileNotFoundError(path)
        s=slide(f'{sid}  ·  {actor}',action)
        # Native image is embedded unchanged; enlarge via the original PNG if needed.
        pic=s.shapes.add_picture(str(path), Inches(.5), Inches(2.15))
        scale=min(15*914400/pic.width, 6.65*914400/pic.height)
        pic.width=int(pic.width*scale);pic.height=int(pic.height*scale)
        pic.left=int((P.slide_width-pic.width)/2)
        status='APROBADA: sólo alcance registrado' if approved else 'APROBADA: lote round-02 completo'
        txt(s,status,.5,8.93,8,.32,13,TEAL,True)
        s.notes_slide.notes_text_frame.text=f'{sid}\n{action}\nFuente: {path.relative_to(ROOT)}\n{status}\nLas flechas navegan la revisión, no ejecutan transacciones. Las condiciones dependen de Odoo.'
        cards.append((sid,actor,action,s));detail_count+=1
    for n,(sid,actor,action,s) in enumerate(cards):
        button(s,'Índice',8,9.03,1.3,index)
        button(s,'Recorrido',9.45,9.03,1.6,intro)
        button(s,'Lámina anterior',11.2,9.03,2,cards[n-1][3] if n else intro)
        button(s,'Lámina siguiente',13.35,9.03,2.15,cards[n+1][3] if n+1<len(cards) else intro)
        # Clickable sequence; descriptions clarify alternatives and conditional steps.
        col=n%2;row=n//2;x=.6+col*7.55;y=2.25+row*1.55
        button(intro,f'{n+1:02d}  {sid}  ·  {actor}',x,y,7.1,s)
        txt(intro,action,x,y+.58,7.1,.8,15,MUTED)
    button(intro,'Volver al índice',12.5,9.03,2.9,index)

for n,(title,s) in enumerate(sections):
    x=.6+(n%2)*7.55;y=2.3+(n//2)*1.5
    button(index,f'{n+1:02d}  {title}',x,y,7.1,s)
    txt(index,f'{len(groups[n][2])} láminas · contexto, actor y siguiente acción',x,y+.63,7,.4,15,MUTED)

end=slide('Qué revisamos y qué todavía no certifica este recorrido','Las imágenes no son nuevas reglas de negocio ni pruebas de implementación.')
items=[('Adaptación','Rejillas: desktop y tablet horizontal. Vertical y teléfono: listas y formularios.'),('Trazabilidad','33 láminas aprobadas de ronda 02 + 3 aprobadas anteriormente; originales preservados.'),('Decisiones','Las flechas indican navegación del storyboard. PIN/Workspace, cartera/venta y factura/toma física son alternativas.'),('Huecos visibles','No están dibujados todos los estados, todas las pestañas ni todos los temas oscuros. No se declaran cerrados.'),('Antes de implementar','Respetar imágenes aprobadas, logo real y contratos de Odoo; la aprobación visual no inventa campos ni permisos.')]
for n,(label,body) in enumerate(items):
    txt(end,label,.7,2.3+n*1.22,3,.5,22,TEAL,True);txt(end,body,4,2.3+n*1.22,11,.8,19)
button(end,'Volver al índice',12,9.03,3.4,index)
button(index,'Límites de la revisión',.6,8.75,4,end)
P.save(OUT/'ORBI_RECORRIDOS_v1.pptx')
print(f'{len(P.slides)} slides, {detail_count} image boards')
